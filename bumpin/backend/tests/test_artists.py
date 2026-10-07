"""Artist side: extraction, checks, highlights, rider and help ticket flows, run sheet.

Runs on the fake LLM provider, so extraction uses the line parser fallback.
"""

from __future__ import annotations

import json
from io import BytesIO
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import db
from backend.app.artists import ai, checks, ripple, routes, tickets  # noqa: F401  (routes registers the reset hook)
from backend.app.artists._compat import AlreadyDecided, decide
from backend.app.artists.runsheet import build_runsheet, runsheet_xlsx
from backend.app.main import app

SPARKLE, HALCYON, NEON, MARLOW, NOVA, DUSK = 1, 2, 3, 4, 5, 6
CLS = SimpleNamespace(confidence=0.95, entity_hint=None, is_major_change=False)
NEON_BODY = (
    "Hey team, pasting Neon Tide's rider below.\n\nTECH\n1x Analogue synth (Moog One)\n"
    "2x CDJ-2000NXS2\n2x wedges\n\nHOSPO\n6x bottled water\n1x fruit platter\n\nCheers, Leo"
)
NOVA_BODY = ("Hi Ravi, bad news: our flight from Sydney has been cancelled. We are rebooked and land at 9:40pm, "
             "so Nova Lane cannot make the 21:15 set. Chris")


@pytest.fixture(autouse=True)
def fresh_db():
    db.reset()


def receive(from_addr: str, body: str = "see attached", attachment: str | None = None, subject: str = "Rider") -> int:
    with db.get_conn() as conn:
        eid = conn.execute(
            "INSERT INTO emails (direction, from_addr, subject, body, received_at) VALUES ('in', ?, ?, ?, ?)",
            (from_addr, subject, body, "2026-11-30T10:00:00"),
        ).lastrowid
        if attachment:
            conn.execute(
                """INSERT INTO documents (owner_type, owner_id, kind, filename, path, received_at, email_id)
                   VALUES ('artist', 0, 'other', ?, ?, ?, ?)""",
                (attachment.rsplit("/", 1)[-1], attachment, "2026-11-30T10:00:00", eid),
            )
    return eid


def sparkle_ticket() -> int:
    return tickets.create_rider_ticket(
        receive("sam@sparkle-mgmt.example.test", attachment="data/docs/riders/sparkle_rider.pdf"), CLS)


def findings(ticket_id: int) -> list[dict]:
    with db.get_conn() as conn:
        return db.rows(conn.execute("SELECT * FROM findings WHERE ticket_id = ?", (ticket_id,)))


def outbox() -> list[dict]:
    with db.get_conn() as conn:
        return db.rows(conn.execute("SELECT * FROM outbox ORDER BY id"))


# --- extraction -------------------------------------------------------------------

def test_heuristic_extract_reads_sections_and_quotes():
    items = ai.heuristic_extract([(1, NEON_BODY)])
    assert [(i.quantity, i.category) for i in items] == [
        (1, "technical"), (2, "technical"), (2, "technical"), (6, "hospitality"), (1, "hospitality")]
    assert items[0].quote == "1x Analogue synth (Moog One)"


def test_heuristic_change_reads_cancellation_and_time():
    change = ai.heuristic_change(NOVA_BODY)
    assert change.kind == "cancellation"
    assert change.earliest_start == "21:40"
    assert "cancelled" in change.quote


# --- checks: one passing and one failing case each ---------------------------------------

def test_shortage_found_for_sparkle():
    tid = sparkle_ticket()
    f = [x for x in findings(tid) if x["kind"] == "shortage"]
    assert len(f) == 1 and f[0]["severity"] == "conflict"
    assert "3x Pioneer CDJ-3000" in f[0]["message"] and "has 2" in f[0]["message"]


def test_no_shortage_for_halcyon():
    assert checks.check_shortage(HALCYON) == []


def test_double_booking_found_for_neon_tide():
    tid = tickets.create_rider_ticket(receive("leo@neontide.example.test", NEON_BODY), CLS)
    f = [x for x in findings(tid) if x["kind"] == "double_booking"]
    assert len(f) == 1 and "Halcyon" in f[0]["message"]


def test_no_double_booking_when_halcyon_rejected():
    with db.get_conn() as conn:
        conn.execute("UPDATE tickets SET status = 'rejected' WHERE owner_id = ?", (HALCYON,))
    tid = tickets.create_rider_ticket(receive("leo@neontide.example.test", NEON_BODY), CLS)
    assert not [x for x in findings(tid) if x["kind"] == "double_booking"]


def test_hospitality_over_budget_for_marlow():
    tid = tickets.create_rider_ticket(
        receive("ada@marlowlanes.example.test", attachment="data/docs/riders/marlow_lanes_rider.pdf"), CLS)
    assert checks.hospitality_total(MARLOW) == 1599
    f = [x for x in findings(tid) if x["kind"] == "over_budget"]
    assert len(f) == 1 and "$1,599" in f[0]["message"] and "$1,200" in f[0]["message"]


def test_hospitality_within_cap_for_halcyon():
    assert checks.hospitality_total(HALCYON) == 8 * 2 + 120
    assert checks.check_hospitality(HALCYON) == []


# --- highlights ------------------------------------------------------------------------

def test_highlight_for_problem_1():
    tid = sparkle_ticket()
    f = next(x for x in findings(tid) if x["kind"] == "shortage")
    with TestClient(app) as client:
        boxes = client.get(f"/api/documents/{f['doc_id']}/highlights").json()
        assert client.get(f"/api/documents/{f['doc_id']}/file").headers["content-type"] == "application/pdf"
    assert len(boxes) == 1
    box = boxes[0]
    assert box["page"] == 2 and box["severity"] == "conflict" and box["finding_id"] == f["id"]
    x0, y0, x1, y1 = box["rect"]
    assert 0 < x0 < x1 < box["page_size"][0] and 0 < y0 < y1 < box["page_size"][1]


# --- rider approve and reject -----------------------------------------------------------

def test_approve_blocked_by_open_conflict_without_override():
    tid = sparkle_ticket()
    with pytest.raises(tickets.RiderBlocked):
        tickets.RiderHandler().approve(tid, "ravi", None)


def test_approve_with_override_allocates_what_exists():
    tid = sparkle_ticket()
    tickets.RiderHandler().approve(tid, "ravi", {"override_reason": "Artist brings a third deck"})
    with db.get_conn() as conn:
        cdj = conn.execute("SELECT quantity FROM allocations WHERE artist_id = ? AND inventory_item_id = 1",
                           (SPARKLE,)).fetchone()[0]
        t = db.row(conn.execute("SELECT status, decided_by FROM tickets WHERE id = ?", (tid,)))
    assert cdj == 2
    assert t == {"status": "approved", "decided_by": "ravi"}


def test_resolve_drafts_reply_then_approve_works():
    tid = sparkle_ticket()
    f = next(x for x in findings(tid) if x["kind"] == "shortage")
    tickets.RiderHandler().resolve_finding(tid, f["id"], "jess")
    drafts = outbox()
    assert len(drafts) == 1 and drafts[0]["status"] == "draft" and drafts[0]["sent_at"] is None
    assert drafts[0]["to_addr"] == "sam@sparkle-mgmt.example.test" and "Hi Sam" in drafts[0]["body"]
    tickets.RiderHandler().approve(tid, "jess", None)


def test_second_decision_raises_already_decided():
    tid = sparkle_ticket()
    tickets.RiderHandler().reject(tid, "jess", "We cannot supply a third deck")
    with pytest.raises(AlreadyDecided) as err:
        decide(tid, "ravi", "approve")
    assert err.value.by == "jess"


def test_reject_drafts_email_and_releases_allocations():
    with db.get_conn() as conn:
        tid = conn.execute("SELECT id FROM tickets WHERE owner_id = ?", (HALCYON,)).fetchone()[0]
        conn.execute("UPDATE tickets SET decided_by = NULL WHERE id = ?", (tid,))
    tickets.reject_rider(tid, "jess", "The synth is needed elsewhere at that time")
    with db.get_conn() as conn:
        reserved = conn.execute("SELECT COUNT(*) FROM allocations WHERE artist_id = ? AND status = 'reserved'",
                                (HALCYON,)).fetchone()[0]
    assert reserved == 0
    draft = outbox()[-1]
    assert draft["subject"].startswith("Update needed") and "rejected" not in draft["subject"].lower()


def test_revised_rider_reopens_ticket():
    tid = sparkle_ticket()
    tickets.reject_rider(tid, "jess", "Three decks is not possible")
    tid2 = sparkle_ticket()
    assert tid2 == tid
    with db.get_conn() as conn:
        t = db.row(conn.execute("SELECT status, decided_by FROM tickets WHERE id = ?", (tid,)))
    assert t == {"status": "needs_review", "decided_by": None}


def test_unknown_sender_goes_to_review():
    tid = tickets.create_rider_ticket(receive("someone@unknown.example.test", NEON_BODY), CLS)
    with db.get_conn() as conn:
        t = db.row(conn.execute("SELECT status, owner_id FROM tickets WHERE id = ?", (tid,)))
    assert t == {"status": "needs_review", "owner_id": None}


# --- help ticket ripple (problem 6) ------------------------------------------------------

def nova_ticket() -> int:
    eid = receive("chris@novalane.example.test", NOVA_BODY, subject="Flight cancelled")
    return ripple.create_help_ticket(eid, SimpleNamespace(confidence=0.95, entity_hint="Nova Lane",
                                                          is_major_change=True))


def test_ripple_proposes_swap_and_notifications():
    tid = nova_ticket()
    with db.get_conn() as conn:
        t = db.row(conn.execute("SELECT * FROM tickets WHERE id = ?", (tid,)))
    actions = json.loads(t["proposed_actions_json"])
    assert t["type"] == "help" and t["status"] == "needs_review" and t["severity"] == "conflict"
    assert [a["kind"] for a in actions] == ["move_set", "move_set", "notify", "notify"]
    assert actions[0]["new_start"] == "2026-12-12T22:45:00"
    assert actions[1]["artist_id"] == DUSK and actions[1]["new_start"] == "2026-12-12T21:15:00"
    assert outbox() == []  # nothing drafted before a human approves


def test_approving_actions_moves_sets_and_drafts_emails():
    tid = nova_ticket()
    for i in range(4):
        ripple.HelpHandler().approve_action(tid, i, "ravi")
    with db.get_conn() as conn:
        sets = {r["id"]: (r["set_start"], r["set_end"]) for r in
                db.rows(conn.execute("SELECT id, set_start, set_end FROM artists WHERE id IN (5, 6)"))}
        status = conn.execute("SELECT status FROM tickets WHERE id = ?", (tid,)).fetchone()[0]
    assert sets[NOVA] == ("2026-12-12T22:45:00", "2026-12-13T00:15:00")
    assert sets[DUSK] == ("2026-12-12T21:15:00", "2026-12-12T22:30:00")
    assert status == "resolved"
    drafts = outbox()
    assert len(drafts) == 4 and all(d["status"] == "draft" for d in drafts)
    with pytest.raises(AlreadyDecided):
        ripple.approve_action(tid, 0, "jess")


def test_edit_action_changes_time():
    tid = nova_ticket()
    action = ripple.HelpHandler()
    action.edit_action(tid, 0, "ravi", {"new_start": "2026-12-12T22:50:00"})
    with db.get_conn() as conn:
        a = json.loads(conn.execute("SELECT proposed_actions_json FROM tickets WHERE id = ?", (tid,)).fetchone()[0])[0]
    assert a["new_start"] == "2026-12-12T22:50:00" and a["new_end"] == "2026-12-13T00:20:00"
    assert a["title"].endswith("22:50")


# --- run sheet, overview and copy rules -----------------------------------------------------

def test_runsheet_and_xlsx_match():
    rows = build_runsheet()
    assert len(rows) == 45 and rows == sorted(rows, key=lambda r: (r.start, r.area))
    ws = load_workbook(BytesIO(runsheet_xlsx())).active
    data = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(data) == len(rows)
    assert [r[4] for r in data] == [r.who for r in rows]


def test_overview_and_inventory_endpoints():
    sparkle_ticket()
    with TestClient(app) as client:
        ov = client.get("/api/overview").json()
        inv = client.get("/api/inventory").json()
        xlsx = client.get("/api/export/runsheet.xlsx")
    assert ov["suppliers"]["artists"] == 60 and ov["suppliers"]["artists_scheduled"] == 45
    assert ov["suppliers"]["artists_applied"] == 15 and ov["open_conflicts"] == 1
    synth = next(i for i in inv if i["id"] == 14)
    assert synth["allocations"][0]["artist_name"] == "Halcyon"
    assert xlsx.status_code == 200 and xlsx.content[:2] == b"PK"


def test_drafts_have_no_em_or_en_dashes():
    tid = sparkle_ticket()
    f = findings(tid)[0]
    tickets.resolve_finding(tid, f["id"], "ravi")
    nova = nova_ticket()
    for i in range(4):
        ripple.approve_action(nova, i, "ravi")
    for d in outbox():
        assert "\u2014" not in d["body"] + d["subject"] and "\u2013" not in d["body"] + d["subject"]
    for f in findings(tid) + findings(nova):
        assert "\u2014" not in f["message"] and "\u2013" not in f["message"]


def test_fifteen_sets_a_day_and_applied_artists_have_no_slot():
    per_day: dict[str, int] = {}
    for r in build_runsheet():
        per_day[r.start[:10]] = per_day.get(r.start[:10], 0) + 1
    assert per_day == {"2026-12-11": 15, "2026-12-12": 15, "2026-12-13": 15}
    with TestClient(app) as client:
        applied = [a for a in client.get("/api/artists").json() if a["status"] == "applied"]
    assert len(applied) == 15 and all(a["set_start"] is None and a["stage_id"] is None for a in applied)
