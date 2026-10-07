"""Run sheet from the database (artist sets and vendor load-ins), plus xlsx export."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from backend.app import db
from backend.app.artists.models import RunsheetRow

HEADERS = ["Day", "Start", "End", "Stage or zone", "Who", "Type", "Status", "Contact"]


def _day(ts: str) -> str:
    return datetime.fromisoformat(ts).strftime("%a %d %b")


def build_runsheet() -> list[RunsheetRow]:
    out: list[RunsheetRow] = []
    with db.get_conn() as conn:
        for a in db.rows(conn.execute(
            """SELECT a.*, s.name AS stage_name FROM artists a LEFT JOIN stages s ON s.id = a.stage_id
               WHERE a.set_start IS NOT NULL""")):
            out.append(RunsheetRow(
                day=_day(a["set_start"]), start=a["set_start"], end=a["set_end"],
                area=a["stage_name"] or "", who=a["name"], kind="set", status=a["status"],
                contact=a["manager_email"],
            ))
        for v in db.rows(conn.execute("SELECT * FROM vendors WHERE load_in_start IS NOT NULL")):
            out.append(RunsheetRow(
                day=_day(v["load_in_start"]), start=v["load_in_start"], end=v["load_in_end"] or v["load_in_start"],
                area=v["site_zone"] or "", who=v["name"], kind="load_in", status=v["status"],
                contact=v["contact_email"],
            ))
    return sorted(out, key=lambda r: (r.start, r.area))


def runsheet_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Run sheet"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E5F")
    for r in build_runsheet():
        ws.append([r.day, r.start[11:16], r.end[11:16], r.area, r.who,
                   "Set" if r.kind == "set" else "Load-in", r.status, r.contact])
    for col, width in zip("ABCDEFGH", (12, 8, 8, 18, 26, 9, 12, 36)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
